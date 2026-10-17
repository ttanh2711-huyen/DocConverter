"""Phương án B: dựng DOCX từ DoclingDocument, cố giữ bố cục trang.

- Khổ trang theo PDF, lề lấy từ khung bao nội dung.
- Các khối nằm cùng "dải" dọc (header 2 bên, trang 2 cột) → bảng không viền.
- Khối đơn: căn giữa / thụt lề theo toạ độ x; cỡ chữ ước lượng từ bbox.
- Bảng giữ ô gộp; ảnh giữ kích thước.
"""

import io
import math
import re
from dataclasses import dataclass
from pathlib import Path

from docling_core.types.doc import DoclingDocument, PictureItem, SectionHeaderItem, TableItem, TextItem
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

FONT = "Times New Roman"
CHAR_W = 0.45  # bề rộng ký tự trung bình / cỡ chữ (Times, tiếng Việt)
LINE_H = 1.15
LINE_TO_FONT = 0.68  # chiều cao dòng OCR (pt) → cỡ chữ; đo trên fixture: 12pt ≈ 16–18
LINE_TO_FONT_UPPER = 0.58  # chữ in hoa có dấu cho hộp dòng cao hơn: 12pt ≈ 20
MIN_MARGIN = 28.0  # pt


@dataclass
class Block:
    item: object
    l: float
    t: float  # toạ độ từ trên xuống (top-left origin)
    r: float
    b: float
    font_size: float | None = None

    @property
    def w(self):
        return self.r - self.l

    @property
    def h(self):
        return self.b - self.t


def _clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def estimate_font_size(text: str, width: float, height: float) -> float:
    """Tìm cỡ chữ sao cho số dòng × chiều cao dòng khớp chiều cao khối."""
    n = max(len(text), 1)
    best, err = 12.0, math.inf
    f = 8.0
    while f <= 30:
        lines = max(1, math.ceil(n * CHAR_W * f / max(width, 1)))
        e = abs(lines * LINE_H * f - height)
        if e < err:
            best, err = f, e
        f += 0.5
    return best


def line_boxes_from(result) -> dict[int, list[tuple[float, float, float]]]:
    """(tâm x, tâm y, chiều cao) của từng dòng text, toạ độ từ trên xuống.

    Cần `PdfPipelineOptions(generate_parsed_pages=True)`; trả về rỗng nếu không có.
    """
    out: dict[int, list[tuple[float, float, float]]] = {}
    for page in result.pages:
        if page.parsed_page is None:
            continue
        page_h = page.size.height
        boxes = []
        for c in page.parsed_page.textline_cells:
            bb = c.rect.to_bounding_box().to_top_left_origin(page_h)
            boxes.append(((bb.l + bb.r) / 2, (bb.t + bb.b) / 2, bb.b - bb.t))
        out[page.page_no + 1] = boxes
    return out


def _font_from_lines(blk: "Block", boxes) -> float | None:
    hs = sorted(h for x, y, h in boxes if blk.l <= x <= blk.r and blk.t <= y <= blk.b)
    if not hs:
        return None
    text = getattr(blk.item, "text", "")
    ratio = LINE_TO_FONT_UPPER if text and text == text.upper() else LINE_TO_FONT
    return min(28.0, max(8.0, round(hs[len(hs) // 2] * ratio * 2) / 2))


def _blocks(doc: DoclingDocument, line_boxes=None):
    line_boxes = line_boxes or {}
    pages: dict[int, list[Block]] = {n: [] for n in doc.pages}
    for item, _ in doc.iterate_items(with_groups=False):
        if not getattr(item, "prov", None):
            continue
        prov = item.prov[0]
        page_h = doc.pages[prov.page_no].size.height
        bb = prov.bbox.to_top_left_origin(page_h)
        blk = Block(item, bb.l, bb.t, bb.r, bb.b)
        blk.font_size = _font_from_lines(blk, line_boxes.get(prov.page_no, []))
        pages[prov.page_no].append(blk)
    return pages


def _bands(blocks: list[Block]) -> list[list[Block]]:
    """Gom các khối chồng nhau theo trục dọc thành một dải; trong dải sắp theo x."""
    bands: list[list[Block]] = []
    for blk in sorted(blocks, key=lambda b: b.t):
        if bands:
            last = bands[-1]
            top, bottom = min(b.t for b in last), max(b.b for b in last)
            overlap = min(bottom, blk.b) - max(top, blk.t)
            side_by_side = all(blk.l >= b.r - 2 or blk.r <= b.l + 2 for b in last)
            if overlap > 0.3 * min(blk.h, bottom - top) and side_by_side:
                last.append(blk)
                continue
        bands.append([blk])
    return [sorted(b, key=lambda x: x.l) for b in bands]


def _set_font(run, size: float, bold: bool = False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)


def _write_text(container, blk: Block, page_w: float, margin_l: float, margin_r: float, in_cell=False):
    item = blk.item
    text = _clean(item.text)
    size = blk.font_size or estimate_font_size(text, blk.w, blk.h)
    bold = isinstance(item, SectionHeaderItem) or item.label == "caption"
    p = container.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(0)
    _set_font(p.add_run(text), size, bold)
    if in_cell:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if bold and blk.w < 0.9 * page_w else WD_ALIGN_PARAGRAPH.JUSTIFY
        return
    body_w = page_w - margin_l - margin_r
    centre_offset = (blk.l + blk.r) / 2 - (margin_l + body_w / 2)
    if blk.w > 0.9 * body_w:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif abs(centre_offset) < 0.05 * body_w and blk.l - margin_l > 10:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.paragraph_format.left_indent = Pt(max(0.0, blk.l - margin_l))


def _write_picture(container, doc: DoclingDocument, blk: Block, margin_l: float, in_cell=False):
    img = blk.item.get_image(doc)
    if img is None:
        return
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    buf.seek(0)
    p = container.add_paragraph()
    if not in_cell:
        p.paragraph_format.left_indent = Pt(max(0.0, blk.l - margin_l))
    p.add_run().add_picture(buf, width=Pt(blk.w), height=Pt(blk.h))


def _write_table(container, blk: Block):
    data = blk.item.data
    table = container.add_table(rows=data.num_rows, cols=data.num_cols)
    _borders(table, "single")
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for c in data.table_cells:
        cell = table.cell(c.start_row_offset_idx, c.start_col_offset_idx)
        end = table.cell(c.end_row_offset_idx - 1, c.end_col_offset_idx - 1)
        if end._tc is not cell._tc:
            cell = cell.merge(end)
        para = cell.paragraphs[0]
        para.text = ""
        _set_font(para.add_run(_clean(c.text)), 11, bold=c.column_header)
        if c.column_header or _clean(c.text).replace(",", "").isdigit():
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def _write_block(container, doc, blk, page_w, ml, mr, in_cell=False):
    if isinstance(blk.item, TableItem):
        _write_table(container, blk)
    elif isinstance(blk.item, PictureItem):
        _write_picture(container, doc, blk, ml, in_cell)
    elif isinstance(blk.item, TextItem):
        _write_text(container, blk, page_w, ml, mr, in_cell)


def _borders(table, val: str):
    """Kẻ khung trực tiếp trên bảng (val="single") hoặc bỏ khung (val="nil")."""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.makeelement(qn("w:tblBorders"), {})
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        attrs = {qn("w:val"): val}
        if val != "nil":
            attrs |= {qn("w:sz"): "6", qn("w:space"): "0", qn("w:color"): "000000"}
        borders.append(borders.makeelement(qn(f"w:{edge}"), attrs))
    tbl_pr.append(borders)


def write_docx(doc: DoclingDocument, out: Path, line_boxes=None) -> Path:
    pages = _blocks(doc, line_boxes)
    d = Document()
    d.styles["Normal"].font.name = FONT
    first = True
    for page_no in sorted(pages):
        blocks = pages[page_no]
        size = doc.pages[page_no].size
        if blocks:
            ml = max(MIN_MARGIN, min(b.l for b in blocks) - 2)
            mr = max(MIN_MARGIN, size.width - max(b.r for b in blocks) - 2)
            mt = max(MIN_MARGIN, min(b.t for b in blocks) - 2)
            # Khoảng trống cuối trang không phải lề: lấy tối đa bằng lề trên.
            mb = max(MIN_MARGIN, min(size.height - max(b.b for b in blocks) - 2, mt))
        else:
            ml = mr = mt = mb = 56.7
        section = d.sections[0] if first else d.add_section(WD_SECTION.NEW_PAGE)
        first = False
        section.page_width, section.page_height = Pt(size.width), Pt(size.height)
        section.left_margin, section.right_margin = Pt(ml), Pt(mr)
        section.top_margin, section.bottom_margin = Pt(mt), Pt(mb)

        for band in _bands(blocks):
            if len(band) == 1:
                _write_block(d, doc, band[0], size.width, ml, mr)
                continue
            grid = d.add_table(rows=1, cols=len(band))
            _borders(grid, "nil")
            for cell, blk in zip(grid.rows[0].cells, band):
                cell.width = Pt(blk.w)
                cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
                _write_block(cell, doc, blk, size.width, ml, mr, in_cell=True)
                if not cell.paragraphs or cell._tc[-1].tag == qn("w:tbl"):
                    cell.add_paragraph()  # Word yêu cầu ô kết thúc bằng một đoạn văn
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(out)
    return out
