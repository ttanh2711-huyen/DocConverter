"""Task 2: so sánh 3 phương án PDF → DOCX trên các PDF mẫu.

A: Docling → HTML → pandoc   B: Docling → docx_writer_proto   C: pdf2docx
Dùng: uv run python pdf_to_docx_compare.py fixtures/*.pdf [--out out/compare]
"""

import argparse
import json
import re
import subprocess
import time
import unicodedata
from pathlib import Path

import pypandoc
import pypdfium2 as pdfium
from docling.datamodel.base_models import InputFormat
from docling.datamodel.pipeline_options import OcrMacOptions, PdfPipelineOptions
from docling.document_converter import DocumentConverter, PdfFormatOption
from docx import Document
from docx.oxml.ns import qn
from pdf2docx import Converter

from docx_writer_proto import line_boxes_from, write_docx

PT_PER_MM = 72 / 25.4


def docling_converter(force_ocr=False):
    opts = PdfPipelineOptions(
        do_ocr=True,
        ocr_options=OcrMacOptions(lang=["vi-VT", "en-US"], force_full_page_ocr=force_ocr),
        generate_picture_images=True,
        generate_parsed_pages=True,
        images_scale=2.0,
    )
    return DocumentConverter(format_options={InputFormat.PDF: PdfFormatOption(pipeline_options=opts)})


def has_text_layer(pdf: Path, min_chars_per_page: int = 20) -> bool:
    doc = pdfium.PdfDocument(str(pdf))
    chars = sum(len(page.get_textpage().get_text_range().strip()) for page in doc)
    return chars >= min_chars_per_page * len(doc)


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", s)).strip().lower()


def pdf_phrases(pdf: Path) -> list[str]:
    """Lấy các dòng dài ≥ 15 ký tự từ lớp text của PDF gốc làm cụm từ kiểm tra."""
    doc = pdfium.PdfDocument(str(pdf))
    lines = []
    for page in doc:
        lines += [norm(x) for x in page.get_textpage().get_text_range().splitlines()]
    return [x for x in dict.fromkeys(lines) if len(x) >= 15]


def analyze(docx_path: Path, pdf: Path, phrases: list[str]) -> dict:
    d = Document(docx_path)
    body = d.element.body
    paras = [" ".join(t.text or "" for t in p.iter(qn("w:t"))) for p in body.iter(qn("w:p"))]
    text = norm(" ".join(paras))
    page = pdfium.PdfDocument(str(pdf))[0]
    pw, ph = page.get_size()
    s = d.sections[0]
    tables = body.findall(".//" + qn("w:tbl"))
    grid_tables = [t for t in tables if not t.findall(".//" + qn("w:tblBorders") + "/" + qn("w:top"))
                   or t.find(".//" + qn("w:tblStyle")) is not None]
    merged = next(body.iter(qn("w:gridSpan"), qn("w:vMerge")), None) is not None
    cols = [c for c in body.iter(qn("w:cols")) if int(c.get(qn("w:num"), "1")) > 1]
    return {
        "page_size_ok": s.page_width is not None
        and abs(s.page_width.pt - pw) < 2 * PT_PER_MM and abs(s.page_height.pt - ph) < 2 * PT_PER_MM,
        "margins_mm": [round(x.pt / PT_PER_MM) if x is not None else None
                       for x in (s.left_margin, s.right_margin, s.top_margin, s.bottom_margin)],
        "sections": len(d.sections),
        "text_hits": f"{sum(p in text for p in phrases)}/{len(phrases)}" if phrases else "-",
        "tables": len(tables),
        "grid_tables": len(grid_tables),
        "merged_cells": merged,
        "images": len(list(body.iter(qn("a:blip")))),
        "column_sections": len(cols),
    }


def thumbnail(docx_path: Path) -> None:
    subprocess.run(["qlmanage", "-t", "-s", "900", "-o", str(docx_path.parent), str(docx_path)],
                   capture_output=True, timeout=60)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pdfs", nargs="+", type=Path)
    ap.add_argument("--out", type=Path, default=Path("out/compare"))
    args = ap.parse_args()

    conv = {False: docling_converter(False), True: docling_converter(True)}
    report = {}
    for pdf in args.pdfs:
        is_text = has_text_layer(pdf)
        phrases = pdf_phrases(pdf) if is_text else pdf_phrases(pdf.with_name(pdf.name.replace("_scan", "")))
        row = {"text_layer": is_text}
        t = time.perf_counter()
        res = conv[not is_text].convert(pdf)
        ddoc = res.document
        t_docling = time.perf_counter() - t

        out_a = args.out / "A" / f"{pdf.stem}.docx"
        out_a.parent.mkdir(parents=True, exist_ok=True)
        t = time.perf_counter()
        pypandoc.convert_text(ddoc.export_to_html(), "docx", format="html", outputfile=str(out_a))
        row["A"] = analyze(out_a, pdf, phrases) | {"seconds": round(t_docling + time.perf_counter() - t, 1)}

        out_b = args.out / "B" / f"{pdf.stem}.docx"
        t = time.perf_counter()
        write_docx(ddoc, out_b, line_boxes_from(res))
        row["B"] = analyze(out_b, pdf, phrases) | {"seconds": round(t_docling + time.perf_counter() - t, 1)}

        if is_text:
            out_c = args.out / "C" / f"{pdf.stem}.docx"
            out_c.parent.mkdir(parents=True, exist_ok=True)
            t = time.perf_counter()
            cv = Converter(str(pdf))
            cv.convert(str(out_c))
            cv.close()
            row["C"] = analyze(out_c, pdf, phrases) | {"seconds": round(time.perf_counter() - t, 1)}

        for k in "ABC":
            p = args.out / k / f"{pdf.stem}.docx"
            if p.exists():
                thumbnail(p)
        report[pdf.name] = row
        print(pdf.name, json.dumps(row, ensure_ascii=False, indent=1))
    (args.out / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=1))


if __name__ == "__main__":
    main()
