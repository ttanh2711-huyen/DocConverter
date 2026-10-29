"""Sổ tổng hợp văn bản (Excel): mỗi văn bản một dòng, cộng dồn qua các lần chuyển đổi.

Tệp đã ghi trước đó (cùng đường dẫn tệp gốc) được cập nhật tại chỗ, không thêm dòng trùng.
Các trang tính khác và những cột người dùng tự thêm phía sau đều được giữ nguyên.
"""

import os
import tempfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from docconv.metadata import DocMeta

SHEET = "Sổ văn bản"
COLUMNS = [  # (tiêu đề, độ rộng)
    ("STT", 6),
    ("Loại văn bản", 16),
    ("Số, ký hiệu", 20),
    ("Ngày ban hành", 14),
    ("Cơ quan ban hành", 32),
    ("Trích yếu nội dung", 50),
    ("Người ký", 22),
    ("Chức vụ người ký", 26),
    ("Tên tệp", 28),
    ("Đường dẫn tệp gốc", 50),
    ("Ngày ghi nhận", 16),
    ("Cần kiểm tra", 24),
]
PATH_COL = 10  # cột khoá: đường dẫn tệp gốc
REQUIRED = [("doc_type", "loại"), ("number", "số"), ("date", "ngày"), ("issuer", "cơ quan"), ("signer", "người ký")]


@dataclass
class Entry:
    src: Path
    meta: DocMeta


@dataclass
class Result:
    added: int
    updated: int
    total: int


def _row(entry: Entry, recorded: datetime) -> list:
    m = entry.meta
    issued = date.fromisoformat(m.date) if m.date else None
    missing = [label for field, label in REQUIRED if not getattr(m, field)]
    return [None, m.doc_type, m.number, issued, m.issuer, m.summary, m.signer, m.position,
            entry.src.name, str(entry.src), recorded, ("Thiếu " + ", ".join(missing)) if missing else ""]


def write(path: Path, entries: list[Entry], now: datetime | None = None) -> Result:
    """Thêm/cập nhật các văn bản vào sổ tại `path` (tạo mới nếu chưa có). Ghi qua tệp tạm rồi thay thế."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    now = (now or datetime.now()).replace(microsecond=0)
    path = path.expanduser()
    if path.exists():
        wb = load_workbook(path)
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.create_sheet(SHEET, 0)
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    warn_fill = PatternFill("solid", fgColor="FFF2CC")
    no_fill = PatternFill(fill_type=None)
    thin = Side(style="thin", color="BFBFBF")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(vertical="top", wrap_text=True)

    if ws.cell(1, 2).value != COLUMNS[1][0]:
        if ws.max_row > 1 or any(cell.value is not None for cell in ws[1]):
            raise ValueError(f"Trang tính \"{SHEET}\" trong {path.name} không đúng mẫu sổ của DocConverter")
        for c, (title, width) in enumerate(COLUMNS, 1):
            cell = ws.cell(1, c, title)
            cell.font, cell.fill, cell.border = header_font, header_fill, border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(c)].width = width
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "C2"

    existing = {ws.cell(r, PATH_COL).value: r for r in range(2, ws.max_row + 1) if ws.cell(r, PATH_COL).value}
    next_row = ws.max_row + 1
    added = updated = 0
    for entry in entries:
        values = _row(entry, now)
        row = existing.get(str(entry.src))
        if row:
            updated += 1
        else:
            row, next_row = next_row, next_row + 1
            existing[str(entry.src)] = row
            added += 1
        for c, value in enumerate(values[1:], 2):
            cell = ws.cell(row, c, value)
            cell.alignment, cell.border = wrap, border
        ws.cell(row, 4).number_format = "DD/MM/YYYY"
        ws.cell(row, 11).number_format = "DD/MM/YYYY HH:MM"
        link = ws.cell(row, 9)
        link.hyperlink = Path(entry.src).as_uri()
        link.font = Font(color="0563C1", underline="single")
        for c in range(1, len(COLUMNS) + 1):
            ws.cell(row, c).fill = warn_fill if values[-1] else no_fill

    last = max(ws.max_row, 1)
    for r in range(2, last + 1):  # đánh lại số thứ tự
        cell = ws.cell(r, 1, r - 1)
        cell.alignment = Alignment(horizontal="center", vertical="top")
        cell.border = border
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last}"

    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".~docconv-", suffix=".xlsx", dir=path.parent)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        if os.path.exists(tmp):
            os.unlink(tmp)
    return Result(added=added, updated=updated, total=last - 1)
