import unicodedata

import pypdfium2 as pdfium
from docx import Document
from docx.oxml.ns import qn

from conftest import run_cli
from docconv.pdf_kind import has_text_layer

MM = 72 / 25.4


def _docx_text(path) -> str:
    body = Document(path).element.body
    paras = (" ".join(t.text or "" for t in p.iter(qn("w:t"))) for p in body.iter(qn("w:p")))
    return unicodedata.normalize("NFC", " ".join(paras))


def _assert_layout(docx_path, pdf_path):
    d = Document(docx_path)
    body = d.element.body
    w, h = pdfium.PdfDocument(str(pdf_path))[0].get_size()
    s = d.sections[0]
    assert abs(s.page_width.pt - w) < 2 * MM and abs(s.page_height.pt - h) < 2 * MM
    assert body.findall(".//" + qn("w:tbl")), "phải có bảng"
    assert next(body.iter(qn("w:gridSpan"), qn("w:vMerge")), None) is not None, "phải giữ ô gộp"
    assert next(body.iter(qn("a:blip")), None) is not None, "phải có ảnh (con dấu)"


def test_has_text_layer():
    from conftest import FIXTURES
    assert has_text_layer(FIXTURES / "layout_vi.pdf")
    assert not has_text_layer(FIXTURES / "layout_scan_vi.pdf")


def test_text_pdf_to_docx_keeps_layout(copy_fixture):
    pdf = copy_fixture("layout_vi.pdf")
    proc = run_cli("convert", pdf, "--to", "docx")
    assert proc.returncode == 0, proc.stderr
    out = pdf.with_suffix(".docx")
    _assert_layout(out, pdf)
    text = _docx_text(out)
    assert "Nguyễn Văn Bình" in text and "Chuyển mục đích sử dụng đất" in text
    assert not pdf.with_suffix(".md").exists()


def test_scan_pdf_to_docx_via_ocr(copy_fixture):
    pdf = copy_fixture("layout_scan_vi.pdf")
    proc = run_cli("convert", pdf, "--to", "docx")
    assert proc.returncode == 0, proc.stderr
    out = pdf.with_suffix(".docx")
    _assert_layout(out, pdf)
    assert "Nguyễn Văn Bình" in _docx_text(out)


def test_md_and_docx_together(copy_fixture):
    pdf = copy_fixture("layout_vi.pdf")
    proc = run_cli("convert", pdf, "--to", "md,docx")
    assert proc.returncode == 0, proc.stderr
    assert pdf.with_suffix(".md").exists() and pdf.with_suffix(".docx").exists()
