"""Sổ tổng hợp Excel: tạo mới, cộng dồn, cập nhật dòng cũ, giữ dữ liệu người dùng; và qua CLI."""

from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

from conftest import events, run_cli
from docconv import register
from docconv.metadata import DocMeta

QD = DocMeta("Quyết định", "1234/QĐ-UBND", "2026-03-15", "ỦY BAN NHÂN DÂN TỈNH CÀ MAU", "Nguyễn Văn Bình",
             "KT. CHỦ TỊCH, PHÓ CHỦ TỊCH", "Về việc phê duyệt kế hoạch")
NOW = datetime(2026, 10, 26, 9, 30)


def rows(path: Path) -> list[tuple]:
    ws = openpyxl.load_workbook(path)[register.SHEET]
    return [tuple(c.value for c in r) for r in ws.iter_rows(min_row=2)]


def test_creates_register_with_header_and_typed_date(tmp_path):
    book = tmp_path / "Sổ tổng hợp.xlsx"
    result = register.write(book, [register.Entry(Path("/a/Quyết định 1234.pdf"), QD)], now=NOW)
    assert (result.added, result.updated, result.total) == (1, 0, 1)
    ws = openpyxl.load_workbook(book)[register.SHEET]
    assert [c.value for c in ws[1]] == [title for title, _ in register.COLUMNS]
    assert ws.freeze_panes == "C2" and ws.auto_filter.ref == "A1:L2"
    row = rows(book)[0]
    assert row[:5] == (1, "Quyết định", "1234/QĐ-UBND", datetime(2026, 3, 15), "ỦY BAN NHÂN DÂN TỈNH CÀ MAU")
    assert row[6:10] == ("Nguyễn Văn Bình", "KT. CHỦ TỊCH, PHÓ CHỦ TỊCH", "Quyết định 1234.pdf", "/a/Quyết định 1234.pdf")
    assert row[11] in ("", None)
    assert ws.cell(2, 4).number_format == "DD/MM/YYYY"


def test_appends_and_updates_same_source_without_duplicates(tmp_path):
    book = tmp_path / "so.xlsx"
    register.write(book, [register.Entry(Path("/a/1.pdf"), QD)], now=NOW)
    partial = DocMeta(doc_type="Công văn", number="45/UBND-KT")
    result = register.write(book, [register.Entry(Path("/a/2.docx"), partial),
                                   register.Entry(Path("/a/1.pdf"), DocMeta(**{**QD.__dict__, "signer": "Lê Thị Hoa"}))],
                            now=NOW)
    assert (result.added, result.updated, result.total) == (1, 1, 2)
    data = rows(book)
    assert [r[0] for r in data] == [1, 2]
    assert data[0][6] == "Lê Thị Hoa"
    assert data[1][11] == "Thiếu ngày, cơ quan, người ký"


def test_keeps_other_sheets_and_user_columns(tmp_path):
    book = tmp_path / "so.xlsx"
    register.write(book, [register.Entry(Path("/a/1.pdf"), QD)], now=NOW)
    wb = openpyxl.load_workbook(book)
    wb[register.SHEET].cell(2, 13, "ghi chú riêng")
    wb.create_sheet("Thống kê")["A1"] = "giữ nguyên"
    wb.save(book)
    register.write(book, [register.Entry(Path("/a/1.pdf"), QD)], now=NOW)
    wb = openpyxl.load_workbook(book)
    assert wb["Thống kê"]["A1"].value == "giữ nguyên"
    assert wb[register.SHEET].cell(2, 13).value == "ghi chú riêng"


def test_refuses_foreign_sheet_with_same_name(tmp_path):
    book = tmp_path / "khac.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = register.SHEET
    wb.active["A1"] = "Dữ liệu của người dùng"
    wb.save(book)
    with pytest.raises(ValueError):
        register.write(book, [register.Entry(Path("/a/1.pdf"), QD)])
    assert openpyxl.load_workbook(book).active["A1"].value == "Dữ liệu của người dùng"


def test_cli_records_documents_and_skips_plain_data(copy_fixture, tmp_path):
    cong_van = copy_fixture("cong_van.docx")
    copy_fixture("thong_ke.xlsx")
    book = tmp_path / "Tổng hợp.xlsx"
    proc = run_cli("convert", cong_van.parent, "--json", "--register", book)
    assert proc.returncode == 0, proc.stderr
    evs = events(proc)
    done = {Path(e["file"]).name: e for e in evs if e["event"] == "done"}
    assert done["cong_van.docx"]["meta"]["number"] == "2468/UBND-KT"
    assert "meta" not in done["thong_ke.xlsx"]
    reg = next(e for e in evs if e["event"] == "register")
    assert (reg["added"], reg["total"]) == (1, 1) and evs[-1]["event"] == "finished"
    row = rows(book)[0]
    assert row[1:4] == ("Công văn", "2468/UBND-KT", datetime(2026, 10, 9))
    assert row[4] == "ỦY BAN NHÂN DÂN THÀNH PHỐ CẦN THƠ"
    assert row[6] == "Trần Việt Trường"


def test_cli_reports_register_error_but_keeps_outputs(copy_fixture, tmp_path):
    cong_van = copy_fixture("cong_van.docx")
    blocker = tmp_path / "là tệp"
    blocker.write_text("x")
    proc = run_cli("convert", cong_van, "--json", "--register", blocker / "so.xlsx")
    evs = events(proc)
    assert proc.returncode == 0
    assert next(e for e in evs if e["event"] == "register")["error"].startswith("Không ghi được sổ tổng hợp")
    assert (cong_van.parent / "cong_van.md").exists()


@pytest.mark.parametrize("to", ["md", "docx"])
def test_cli_reads_two_column_pdf(copy_fixture, tmp_path, to):
    """Cả hai đường: Markdown từ Docling (md) và lớp chữ của PDF (chỉ xuất Word)."""
    pdf = copy_fixture("quyet_dinh_2cot.pdf")
    proc = run_cli("convert", pdf, "--json", "--to", to, "--register", tmp_path / "so.xlsx")
    assert proc.returncode == 0, proc.stderr
    meta = next(e for e in events(proc) if e["event"] == "done")["meta"]
    assert meta == {
        "doc_type": "Quyết định", "number": "1234/QĐ-UBND", "date": "2026-03-15",
        "issuer": "ỦY BAN NHÂN DÂN TỈNH CÀ MAU", "signer": "Nguyễn Văn Bình",
        "position": "KT. CHỦ TỊCH, PHÓ CHỦ TỊCH",
        "summary": "Về việc phê duyệt Kế hoạch sử dụng đất năm 2026 của huyện Trần Văn Thời",
    }
    assert rows(tmp_path / "so.xlsx")[0][11] in ("", None)
