"""Tạo các tệp mẫu không phải PDF cho test markitdown: uv run python tests/fixtures/make_office.py"""

from pathlib import Path

import openpyxl
from docx import Document
from pptx import Presentation

HERE = Path(__file__).parent


def make_docx(path: Path) -> None:
    d = Document()
    d.add_heading("Quyết định giao đất", level=1)
    d.add_paragraph("Ủy ban nhân dân thành phố Cần Thơ giao đất cho hộ ông Trần Văn Lộc.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Thửa số", "Diện tích (m²)"
    t.cell(1, 0).text, t.cell(1, 1).text = "152", "1.250,5"
    d.save(path)


def make_cong_van(path: Path) -> None:
    """Văn bản hành chính đúng thể thức: phần đầu và phần ký trong bảng hai cột không viền."""
    d = Document()
    head = d.add_table(rows=2, cols=2)
    head.cell(0, 0).text = "ỦY BAN NHÂN DÂN\nTHÀNH PHỐ CẦN THƠ"
    head.cell(0, 1).text = "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc"
    head.cell(1, 0).text = "Số: 2468/UBND-KT\nV/v rà soát quỹ đất công năm 2027"
    head.cell(1, 1).text = "Cần Thơ, ngày 09 tháng 10 năm 2026"
    d.add_paragraph("Kính gửi: Ủy ban nhân dân các quận, huyện.")
    d.add_paragraph("Ủy ban nhân dân thành phố đề nghị các đơn vị rà soát quỹ đất công trên địa bàn.")
    sign = d.add_table(rows=1, cols=2)
    sign.cell(0, 0).text = "Nơi nhận:\n- Như trên;\n- Lưu: VT."
    sign.cell(0, 1).text = "KT. CHỦ TỊCH\nPHÓ CHỦ TỊCH\n(Đã ký)\nTrần Việt Trường"
    d.save(path)


def make_xlsx(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Thống kê"
    for row in [("Phường", "Số hồ sơ"), ("Tân An", 120), ("Xuân Khánh", 98)]:
        ws.append(row)
    wb.save(path)


def make_pptx(path: Path) -> None:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Kế hoạch sử dụng đất năm 2027"
    slide.placeholders[1].text = "Rà soát quỹ đất công trên địa bàn quận"
    prs.save(path)


if __name__ == "__main__":
    make_docx(HERE / "quyet_dinh.docx")
    make_cong_van(HERE / "cong_van.docx")
    make_xlsx(HERE / "thong_ke.xlsx")
    make_pptx(HERE / "ke_hoach.pptx")
    (HERE / "thong_bao.html").write_text(
        "<html><head><meta charset='utf-8'><title>Thông báo</title></head><body>"
        "<h1>Thông báo lịch tiếp công dân</h1><p>Thứ Năm hằng tuần tại trụ sở phường.</p></body></html>",
        encoding="utf-8")
    (HERE / "danh_sach.csv").write_text("Họ tên,Thửa đất\nLê Thị Hồng,87\nPhạm Quốc Việt,91\n", encoding="utf-8")
    (HERE / "cu.doc").write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 504)  # chỉ là header OLE, không hỗ trợ
    print("Đã tạo các tệp mẫu Office/HTML/CSV")
